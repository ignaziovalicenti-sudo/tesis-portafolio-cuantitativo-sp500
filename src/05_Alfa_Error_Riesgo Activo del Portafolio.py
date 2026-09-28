

# ==============================================================
# SCRIPT FINAL
# ALFA POR REGRESION, TRACKING ERROR Y RIESGO ACTIVO
#
# Archivos admitidos 
# 1. serie_diaria_portafolio_spmo.csv
# 2. reconstruccion_portafolio_spmo.xlsx
#
# Ambos deben haber sido generados por el script unificado
# "04_Rendimiento Portafolio vs. SPMO".
# ==============================================================

!pip -q install statsmodels openpyxl

import warnings
warnings.filterwarnings("ignore")

import os
import numpy as np
import pandas as pd
import statsmodels.api as sm
import matplotlib.pyplot as plt

from google.colab import files
from IPython.display import display
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
# 1. PARAMETROS
# ============================================================

CAPITAL_INICIAL = 100000.00

BENCHMARK = "SPMO"

DIAS_BURSATILES = 252

# Tasa anual decimal:
# 5,18% = 0.0518
TASA_RF_ANUAL = 0.0518

NIVEL_SIGNIFICANCIA = 0.05

ARCHIVO_EXCEL_SALIDA = (
    "resultados_alfa_tracking_error.xlsx"
)

ARCHIVO_GRAFICO = (
    "regresion_alfa_portafolio.png"
)


# ============================================================
# 2. CARGAR ARCHIVO DE SERIE DIARIA
# ============================================================

print(
    "Seleccione uno de estos archivos:\n"
    "1. serie_diaria_portafolio_spmo.csv\n"
    "2. reconstruccion_portafolio_spmo.xlsx"
)

archivos_subidos = files.upload()

if len(archivos_subidos) == 0:

    raise ValueError(
        "No se seleccionó ningún archivo."
    )

archivo_seleccionado = next(
    iter(archivos_subidos)
)

extension = os.path.splitext(
    archivo_seleccionado
)[1].lower()

print(
    f"\nArchivo seleccionado: "
    f"{archivo_seleccionado}"
)


# ============================================================
# 3. LEER CSV O EXCEL
# ============================================================

if extension == ".csv":

    df = pd.read_csv(
        archivo_seleccionado
    )

elif extension in [".xlsx", ".xls"]:

    libro_excel = pd.ExcelFile(
        archivo_seleccionado
    )

    print(
        "Hojas encontradas:",
        libro_excel.sheet_names
    )

    if "Serie diaria" not in libro_excel.sheet_names:

        raise ValueError(
            "El Excel no contiene la hoja "
            "'Serie diaria'.\n"
            "Seleccione el archivo "
            "'reconstruccion_portafolio_spmo.xlsx'."
        )

    df = pd.read_excel(
        archivo_seleccionado,
        sheet_name="Serie diaria"
    )

else:

    raise ValueError(
        "Formato no admitido. "
        "Seleccione un archivo CSV o XLSX."
    )


# ============================================================
# 4. NORMALIZAR NOMBRES DE COLUMNAS
# ============================================================

df.columns = (
    df.columns
    .astype(str)
    .str.strip()
)

# El índice exportado podría llamarse Fecha, Date
# o aparecer como columna sin nombre.

if "Fecha" not in df.columns:

    if "Date" in df.columns:

        df = df.rename(
            columns={"Date": "Fecha"}
        )

    elif "Unnamed: 0" in df.columns:

        df = df.rename(
            columns={"Unnamed: 0": "Fecha"}
        )


print("\nColumnas encontradas:")

print(df.columns.tolist())

display(df.head())


# ============================================================
# 5. DETECTAR ARCHIVO INCORRECTO
# ============================================================

columnas_operaciones = {
    "Symbol",
    "TransactionType",
    "Quantity",
    "Price"
}

if columnas_operaciones.issubset(
    set(df.columns)
):

    raise ValueError(
        "\nARCHIVO INCORRECTO\n\n"
        "Seleccionó el historial de operaciones de "
        "Wall Street Survivor.\n\n"
        "Este script necesita el archivo generado por "
        "el script del gráfico:\n\n"
        "serie_diaria_portafolio_spmo.csv\n\n"
        "o bien:\n\n"
        "reconstruccion_portafolio_spmo.xlsx"
    )


# ============================================================
# 6. VALIDAR COLUMNAS NECESARIAS
# ============================================================

columnas_requeridas = [
    "Fecha",
    "Cash",
    "Assets",
    "PortfolioValue",
    BENCHMARK
]

columnas_faltantes = [
    columna
    for columna in columnas_requeridas
    if columna not in df.columns
]

if columnas_faltantes:

    raise ValueError(
        "Faltan las siguientes columnas:\n"
        f"{columnas_faltantes}\n\n"
        "Columnas encontradas:\n"
        f"{df.columns.tolist()}\n\n"
        "Seleccione el archivo de serie diaria "
        "generado por el script del gráfico."
    )


# ============================================================
# 7. LIMPIEZA DE LA SERIE
# ============================================================

df["Fecha"] = pd.to_datetime(
    df["Fecha"],
    errors="coerce"
)

columnas_numericas = [
    "Cash",
    "Assets",
    "PortfolioValue",
    BENCHMARK
]

for columna in columnas_numericas:

    df[columna] = pd.to_numeric(
        df[columna],
        errors="coerce"
    )

df = (
    df
    .dropna(
        subset=[
            "Fecha",
            "Cash",
            "Assets",
            "PortfolioValue",
            BENCHMARK
        ]
    )
    .sort_values("Fecha")
    .drop_duplicates(
        subset=["Fecha"],
        keep="last"
    )
    .set_index("Fecha")
)

if len(df) < 10:

    raise ValueError(
        "La serie contiene menos de 10 observaciones."
    )

print("\nPeríodo analizado:")

print(
    "Desde:",
    df.index.min().date()
)

print(
    "Hasta:",
    df.index.max().date()
)

print(
    "Cantidad de fechas:",
    len(df)
)

display(df.tail())


# ============================================================
# 8. CONTROL DEL PATRIMONIO
# ============================================================

df["Control_Patrimonio"] = (
    df["Cash"]
    + df["Assets"]
)

df["Diferencia_Control"] = (
    df["PortfolioValue"]
    - df["Control_Patrimonio"]
)

diferencia_maxima_control = (
    df["Diferencia_Control"]
    .abs()
    .max()
)

print(
    "\nDiferencia máxima entre PortfolioValue "
    "y Cash + Assets:"
)

print(
    f"USD {diferencia_maxima_control:,.6f}"
)

if diferencia_maxima_control <= 0.01:

    print(
        "Control correcto: el patrimonio coincide."
    )

else:

    print(
        "Advertencia: existe una diferencia superior "
        "a USD 0,01."
    )


# ============================================================
# 9. PERFORMANCE ACUMULADA
# ============================================================

patrimonio_final = float(
    df["PortfolioValue"].iloc[-1]
)

cash_final = float(
    df["Cash"].iloc[-1]
)

activos_finales = float(
    df["Assets"].iloc[-1]
)

rendimiento_portafolio = (
    patrimonio_final
    / CAPITAL_INICIAL
    - 1
)

rendimiento_spmo = (
    df[BENCHMARK].iloc[-1]
    / df[BENCHMARK].iloc[0]
    - 1
)

diferencia_acumulada = (
    rendimiento_portafolio
    - rendimiento_spmo
)


# ============================================================
# 10. RENDIMIENTOS DIARIOS
# ============================================================

datos = df[
    [
        "PortfolioValue",
        BENCHMARK
    ]
].copy()

datos["Ret_Portafolio"] = (
    datos["PortfolioValue"]
    .pct_change(fill_method=None)
)

datos["Ret_Benchmark"] = (
    datos[BENCHMARK]
    .pct_change(fill_method=None)
)

datos = (
    datos
    .replace(
        [np.inf, -np.inf],
        np.nan
    )
    .dropna(
        subset=[
            "Ret_Portafolio",
            "Ret_Benchmark"
        ]
    )
)

if len(datos) < 10:

    raise ValueError(
        "No existen suficientes rendimientos diarios "
        "para estimar la regresión."
    )


# ============================================================
# 11. TASA LIBRE DE RIESGO
# ============================================================

tasa_rf_diaria = (
    (1 + TASA_RF_ANUAL)
    ** (1 / DIAS_BURSATILES)
    - 1
)

datos["Rf_Diaria"] = tasa_rf_diaria

datos["Exceso_Portafolio"] = (
    datos["Ret_Portafolio"]
    - datos["Rf_Diaria"]
)

datos["Exceso_Benchmark"] = (
    datos["Ret_Benchmark"]
    - datos["Rf_Diaria"]
)


# ============================================================
# 12. REGRESION DEL ALFA
#
# Rp - Rf = alfa + beta x (Rb - Rf) + error
# ============================================================

Y = datos["Exceso_Portafolio"]

X = sm.add_constant(
    datos["Exceso_Benchmark"],
    has_constant="add"
)

modelo = sm.OLS(
    Y,
    X,
    missing="drop"
).fit()

alfa_diario = float(
    modelo.params["const"]
)

beta = float(
    modelo.params["Exceso_Benchmark"]
)

error_estandar_alfa = float(
    modelo.bse["const"]
)

estadistico_t_alfa = float(
    modelo.tvalues["const"]
)

p_valor_alfa = float(
    modelo.pvalues["const"]
)

r_cuadrado = float(
    modelo.rsquared
)

r_cuadrado_ajustado = float(
    modelo.rsquared_adj
)

cantidad_observaciones = int(
    modelo.nobs
)


# ============================================================
# 13. INTERVALO DE CONFIANZA DEL ALFA
# ============================================================

intervalo_confianza = modelo.conf_int(
    alpha=NIVEL_SIGNIFICANCIA
)

ic_alfa_diario_inferior = float(
    intervalo_confianza.loc["const", 0]
)

ic_alfa_diario_superior = float(
    intervalo_confianza.loc["const", 1]
)


# ============================================================
# 14. ANUALIZACION DEL ALFA
# ============================================================

alfa_anualizado = (
    (1 + alfa_diario)
    ** DIAS_BURSATILES
    - 1
)

ic_alfa_anual_inferior = (
    (1 + ic_alfa_diario_inferior)
    ** DIAS_BURSATILES
    - 1
)

ic_alfa_anual_superior = (
    (1 + ic_alfa_diario_superior)
    ** DIAS_BURSATILES
    - 1
)


# ============================================================
# 15. TRACKING ERROR Y RIESGO ACTIVO
# ============================================================

datos["Rendimiento_Activo"] = (
    datos["Ret_Portafolio"]
    - datos["Ret_Benchmark"]
)

tracking_error_diario = float(
    datos["Rendimiento_Activo"]
    .std(ddof=1)
)

tracking_error_anualizado = (
    tracking_error_diario
    * np.sqrt(DIAS_BURSATILES)
)

# Tracking error y riesgo activo son la misma medida.

riesgo_activo_anualizado = (
    tracking_error_anualizado
)


# ============================================================
# 16. RENDIMIENTO ACTIVO E INFORMATION RATIO
# ============================================================

rendimiento_activo_medio_diario = float(
    datos["Rendimiento_Activo"]
    .mean()
)

rendimiento_activo_anualizado = (
    rendimiento_activo_medio_diario
    * DIAS_BURSATILES
)

if tracking_error_anualizado > 0:

    information_ratio = (
        rendimiento_activo_anualizado
        / tracking_error_anualizado
    )

else:

    information_ratio = np.nan


# ============================================================
# 17. SIGNIFICANCIA ESTADISTICA
# ============================================================

if p_valor_alfa < 0.01:

    significancia = "Significativo al 1%"

elif p_valor_alfa < 0.05:

    significancia = "Significativo al 5%"

elif p_valor_alfa < 0.10:

    significancia = "Significativo al 10%"

else:

    significancia = (
        "No estadísticamente significativo"
    )

if alfa_diario > 0:

    signo_alfa = "Positivo"

elif alfa_diario < 0:

    signo_alfa = "Negativo"

else:

    signo_alfa = "Igual a cero"


# ============================================================
# 18. TABLA DE RESULTADOS
# ============================================================

resultados = pd.DataFrame({
    "Indicador": [
        "Capital inicial",
        "Patrimonio final",
        "Cash final",
        "Valor posiciones abiertas",
        "Rendimiento acumulado del portafolio",
        "Rendimiento acumulado de SPMO",
        "Diferencia acumulada",
        "Tasa libre de riesgo anual",
        "Tasa libre de riesgo diaria",
        "Alfa diario",
        "Alfa anualizado",
        "IC 95% alfa anual inferior",
        "IC 95% alfa anual superior",
        "Error estándar del alfa",
        "Estadístico t del alfa",
        "Valor p del alfa",
        "Beta",
        "R cuadrado",
        "R cuadrado ajustado",
        "Tracking error diario",
        "Tracking error anualizado",
        "Riesgo activo anualizado",
        "Rendimiento activo anualizado",
        "Information ratio",
        "Cantidad de observaciones",
        "Signo del alfa",
        "Significancia del alfa"
    ],
    "Valor": [
        CAPITAL_INICIAL,
        patrimonio_final,
        cash_final,
        activos_finales,
        rendimiento_portafolio,
        rendimiento_spmo,
        diferencia_acumulada,
        TASA_RF_ANUAL,
        tasa_rf_diaria,
        alfa_diario,
        alfa_anualizado,
        ic_alfa_anual_inferior,
        ic_alfa_anual_superior,
        error_estandar_alfa,
        estadistico_t_alfa,
        p_valor_alfa,
        beta,
        r_cuadrado,
        r_cuadrado_ajustado,
        tracking_error_diario,
        tracking_error_anualizado,
        riesgo_activo_anualizado,
        rendimiento_activo_anualizado,
        information_ratio,
        cantidad_observaciones,
        signo_alfa,
        significancia
    ]
})


# ============================================================
# 19. MOSTRAR RESULTADOS
# ============================================================

print("\n" + "=" * 68)
print("PERFORMANCE ACUMULADA")
print("=" * 68)

print(
    f"Patrimonio final:          "
    f"USD {patrimonio_final:,.2f}"
)

print(
    f"Rendimiento portafolio:    "
    f"{rendimiento_portafolio:.2%}"
)

print(
    f"Rendimiento SPMO:          "
    f"{rendimiento_spmo:.2%}"
)

print(
    f"Diferencia acumulada:      "
    f"{diferencia_acumulada:.2%}"
)

print("\n" + "=" * 68)
print("ALFA POR REGRESION")
print("=" * 68)

print(
    f"Alfa diario:               "
    f"{alfa_diario:.6%}"
)

print(
    f"Alfa anualizado:           "
    f"{alfa_anualizado:.2%}"
)

print(
    f"Beta:                      "
    f"{beta:.4f}"
)

print(
    f"Estadístico t:             "
    f"{estadistico_t_alfa:.4f}"
)

print(
    f"Valor p:                   "
    f"{p_valor_alfa:.6f}"
)

print(
    f"R cuadrado:                "
    f"{r_cuadrado:.4f}"
)

print(
    f"Significancia:             "
    f"{significancia}"
)

print("\n" + "=" * 68)
print("TRACKING ERROR Y RIESGO ACTIVO")
print("=" * 68)

print(
    f"Tracking error diario:     "
    f"{tracking_error_diario:.4%}"
)

print(
    f"Tracking error anualizado: "
    f"{tracking_error_anualizado:.2%}"
)

print(
    f"Riesgo activo anualizado:  "
    f"{riesgo_activo_anualizado:.2%}"
)

print(
    f"Rendimiento activo anual:  "
    f"{rendimiento_activo_anualizado:.2%}"
)

print(
    f"Information ratio:         "
    f"{information_ratio:.4f}"
)

print("\nRESUMEN ESTADISTICO")

print(modelo.summary())


# ============================================================
# 20. GRAFICO DE REGRESION
# ============================================================

fig, ax = plt.subplots(
    figsize=(12, 8)
)

ax.scatter(
    datos["Exceso_Benchmark"] * 100,
    datos["Exceso_Portafolio"] * 100,
    color="blue",
    alpha=0.70,
    s=55,
    label="Rendimientos diarios"
)

x_linea = np.linspace(
    datos["Exceso_Benchmark"].min(),
    datos["Exceso_Benchmark"].max(),
    100
)

y_linea = (
    alfa_diario
    + beta * x_linea
)

ax.plot(
    x_linea * 100,
    y_linea * 100,
    color="black",
    linewidth=2.5,
    label="Recta de regresión"
)

ax.axhline(
    y=0,
    color="gray",
    linestyle="--",
    linewidth=1
)

ax.axvline(
    x=0,
    color="gray",
    linestyle="--",
    linewidth=1
)

ax.set_title(
    "Regresión de Rendimientos Diarios: "
    "Portafolio vs. SPMO",
    fontsize=18,
    fontweight="bold",
    pad=20
)

ax.set_xlabel(
    "Exceso de rendimiento diario de SPMO (%)",
    fontsize=14,
    labelpad=12
)

ax.set_ylabel(
    "Exceso de rendimiento diario del portafolio (%)",
    fontsize=14,
    labelpad=12
)

ax.tick_params(
    axis="both",
    labelsize=12
)

ax.grid(
    True,
    alpha=0.30
)

ax.legend(
    fontsize=13
)

texto_resultados = (
    f"Alfa diario: {alfa_diario:.4%}\n"
    f"Beta: {beta:.3f}\n"
    f"p-valor: {p_valor_alfa:.4f}\n"
    f"R²: {r_cuadrado:.3f}"
)

ax.text(
    0.03,
    0.97,
    texto_resultados,
    transform=ax.transAxes,
    fontsize=12,
    verticalalignment="top",
    bbox={
        "boxstyle": "round",
        "facecolor": "white",
        "alpha": 0.85
    }
)

plt.tight_layout()

plt.savefig(
    ARCHIVO_GRAFICO,
    dpi=600,
    bbox_inches="tight"
)

plt.show()


# ============================================================
# 21. EXPORTAR A EXCEL
# ============================================================

coeficientes = pd.DataFrame({
    "Coeficiente": modelo.params,
    "Error estándar": modelo.bse,
    "Estadístico t": modelo.tvalues,
    "Valor p": modelo.pvalues
})

intervalos = modelo.conf_int(
    alpha=NIVEL_SIGNIFICANCIA
)

intervalos.columns = [
    "Límite inferior",
    "Límite superior"
]

with pd.ExcelWriter(
    ARCHIVO_EXCEL_SALIDA,
    engine="openpyxl"
) as writer:

    resultados.to_excel(
        writer,
        sheet_name="Resumen",
        index=False
    )

    datos.to_excel(
        writer,
        sheet_name="Rendimientos diarios"
    )

    coeficientes.to_excel(
        writer,
        sheet_name="Regresion"
    )

    intervalos.to_excel(
        writer,
        sheet_name="Intervalos confianza"
    )

    df.to_excel(
        writer,
        sheet_name="Serie original"
    )


# ============================================================
# 22. FORMATO DEL EXCEL
# ============================================================

libro = load_workbook(
    ARCHIVO_EXCEL_SALIDA
)

relleno_encabezado = PatternFill(
    fill_type="solid",
    fgColor="1F4E78"
)

fuente_encabezado = Font(
    color="FFFFFF",
    bold=True
)

for hoja in libro.worksheets:

    hoja.freeze_panes = "A2"

    # Aplicar formato a la primera fila
    for celda in hoja[1]:

        celda.fill = relleno_encabezado
        celda.font = fuente_encabezado.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # Ajustar ancho de columnas
    for columna in hoja.columns:

        ancho_maximo = max(
            len(str(celda.value))
            if celda.value is not None
            else 0
            for celda in columna
        )

        letra_columna = columna[0].column_letter

        hoja.column_dimensions[
            letra_columna
        ].width = min(
            ancho_maximo + 2,
            38
        )

libro.save(
    ARCHIVO_EXCEL_SALIDA
)

print(
    f"Archivo Excel formateado correctamente: "
    f"{ARCHIVO_EXCEL_SALIDA}"
)
